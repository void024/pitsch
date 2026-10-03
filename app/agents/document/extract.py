"""Deterministic text extraction from pitch documents. No LLM involved.

Supported: PDF, PowerPoint (.pptx), plain text/markdown, or text the backend already extracted.
Scanned (image-only) PDFs yield no text and are reported as unreadable rather than guessed at.
"""

import base64
import binascii
import io
import logging
from dataclasses import dataclass, field

from app.agents.document.schemas import DocumentInput

logger = logging.getLogger(__name__)

PDF_TYPES = {"application/pdf"}
PPTX_TYPES = {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
TEXT_TYPES = {"text/plain", "text/markdown", "text/csv"}


@dataclass
class ExtractedDoc:
    document_id: str
    ref: str
    filename: str
    pages: list[str] = field(default_factory=list)
    warning: str | None = None

    @property
    def readable(self) -> bool:
        return sum(len(p.strip()) for p in self.pages) >= 20

    @property
    def chars(self) -> int:
        return sum(len(p) for p in self.pages)


def _kind(doc: DocumentInput, raw: bytes | None) -> str:
    name = doc.filename.lower()
    if doc.mime_type in PDF_TYPES or name.endswith(".pdf") or (raw or b"")[:5] == b"%PDF-":
        return "pdf"
    if doc.mime_type in PPTX_TYPES or name.endswith(".pptx"):
        return "pptx"
    if doc.mime_type in TEXT_TYPES or name.endswith((".txt", ".md", ".csv")):
        return "text"
    return "unknown"


def _pdf_pages(raw: bytes) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # many decks are "encrypted" with an empty password
        except Exception as exc:
            raise ValueError("PDF is password-protected") from exc
    return [(page.extract_text() or "") for page in reader.pages]


def _pptx_pages(raw: bytes) -> list[str]:
    from pptx import Presentation

    pages = []
    for slide in Presentation(io.BytesIO(raw)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        pages.append("\n".join(p for p in parts if p.strip()))
    return pages


def extract(doc: DocumentInput, ref: str, max_bytes: int) -> ExtractedDoc:
    out = ExtractedDoc(document_id=doc.document_id, ref=ref, filename=doc.filename)

    if doc.text is not None and doc.text.strip():
        # Backend-extracted text; form feeds (\f) mark page breaks if present.
        out.pages = doc.text.split("\f")
        return out

    if not doc.content_base64:
        out.warning = "No text or file content was provided."
        return out

    try:
        raw = base64.b64decode(doc.content_base64, validate=True)
    except (binascii.Error, ValueError):
        out.warning = "File content is not valid base64."
        return out
    if len(raw) > max_bytes:
        out.warning = f"File is larger than the {max_bytes // (1024 * 1024)} MB limit."
        return out

    kind = _kind(doc, raw)
    try:
        if kind == "pdf":
            out.pages = _pdf_pages(raw)
        elif kind == "pptx":
            out.pages = _pptx_pages(raw)
        elif kind == "text":
            out.pages = raw.decode("utf-8", errors="replace").split("\f")
        else:
            out.warning = f"Unsupported file type ({doc.mime_type}). Convert to PDF in the backend."
            return out
    except Exception as exc:  # corrupt files must not crash the agent
        logger.warning("document extraction failed", extra={"event": "extract_failed"})
        out.warning = f"Could not read file: {exc.__class__.__name__}."
        out.pages = []
        return out

    if not out.readable:
        out.warning = ("No extractable text found — the file may be scanned images. "
                       "OCR is not enabled.")
    return out
